"""Export a SOAP note to PDF or DOCX for the EHR sign-off workflow.

PDF generation uses `fpdf2` (pure Python, zero system dependencies)
instead of WeasyPrint. WeasyPrint needs the GTK/Pango/Cairo native
libraries, which are not installed on a fresh Windows machine and are a
common source of import-time crashes — the exact kind of Windows
dependency trap this project has hit before (see PyAV/faster-whisper in
the README). fpdf2 avoids that entirely.
"""

import os
import tempfile

from docx import Document
from fpdf import FPDF


def _get_field(soap_note, key):
    if not isinstance(soap_note, dict):
        return str(soap_note) if key == "assessment" else ""
    return soap_note.get(key) or soap_note.get(key.capitalize(), "")


def generate_pdf_from_soap(soap_note):
    subjective = _get_field(soap_note, "subjective")
    objective = _get_field(soap_note, "objective")
    assessment = _get_field(soap_note, "assessment")
    plan = _get_field(soap_note, "plan")

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "SOAP Note", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    for title, content in (
        ("Subjective", subjective),
        ("Objective", objective),
        ("Assessment", assessment),
        ("Plan", plan),
    ):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, (content or "Not documented.").encode("latin-1", "replace").decode("latin-1"))
        pdf.ln(2)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp.close()
    pdf.output(tmp.name)
    return tmp.name


def generate_docx_from_soap(soap_note):
    doc = Document()
    doc.add_heading('SOAP Note', level=1)

    def add_section(title, content):
        doc.add_heading(title, level=2)
        doc.add_paragraph(str(content) if content else 'Not documented.')

    add_section('Subjective', _get_field(soap_note, 'subjective'))
    add_section('Objective', _get_field(soap_note, 'objective'))
    add_section('Assessment', _get_field(soap_note, 'assessment'))
    add_section('Plan', _get_field(soap_note, 'plan'))

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


def remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass
