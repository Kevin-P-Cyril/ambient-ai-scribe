from weasyprint import HTML
from docx import Document
import tempfile
import os


def render_soap_html(soap_note):
    # soap_note may be a dict or string; render a simple HTML template
    if isinstance(soap_note, dict):
        assessment = soap_note.get('Assessment') or soap_note.get('assessment', '')
        plan = soap_note.get('Plan', '')
        subjective = soap_note.get('Subjective', '')
        objective = soap_note.get('Objective', '')
    else:
        assessment = str(soap_note)
        plan = subjective = objective = ''

    html = f"""
    <!doctype html>
    <html>
      <head>
        <meta charset="utf-8" />
        <style>
          body {{ font-family: Arial, Helvetica, sans-serif; margin: 24px; color: #111; }}
          h1 {{ color: #1f2937; font-size: 20px; }}
          h2 {{ color: #374151; font-size: 16px; margin-top: 16px; }}
          pre {{ background:#f9fafb; padding:12px; border-radius:6px; white-space:pre-wrap }}
        </style>
      </head>
      <body>
        <h1>SOAP Note</h1>
        <h2>Subjective</h2>
        <pre>{subjective}</pre>
        <h2>Objective</h2>
        <pre>{objective}</pre>
        <h2>Assessment</h2>
        <pre>{assessment}</pre>
        <h2>Plan</h2>
        <pre>{plan}</pre>
      </body>
    </html>
    """
    return html


def generate_pdf_from_soap(soap_note):
    html = render_soap_html(soap_note)
    with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp:
        HTML(string=html).write_pdf(tmp.name)
        return tmp.name


def generate_docx_from_soap(soap_note):
    doc = Document()
    doc.add_heading('SOAP Note', level=1)

    def add_section(title, content):
        doc.add_heading(title, level=2)
        if content:
            doc.add_paragraph(str(content))
        else:
            doc.add_paragraph('')

    if isinstance(soap_note, dict):
        add_section('Subjective', soap_note.get('Subjective', ''))
        add_section('Objective', soap_note.get('Objective', ''))
        add_section('Assessment', soap_note.get('Assessment', soap_note.get('assessment', '')))
        add_section('Plan', soap_note.get('Plan', ''))
    else:
        add_section('Assessment', str(soap_note))

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.docx')
    tmp.close()
    doc.save(tmp.name)
    return tmp.name


def remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass
